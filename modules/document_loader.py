from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class SourceDocument:
    document_id: str
    title: str
    source: str
    source_type: str = "unknown"
    publication_date: Optional[str] = None
    credibility_info: Dict[str, Any] = field(default_factory=dict)
    original_url: str = ""
    domain: str = "unknown"
    text: str = ""
    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_metadata(self) -> Dict[str, Any]:
        return {
            "document_id": self.document_id,
            "title": self.title,
            "source": self.source,
            "source_type": self.source_type,
            "publication_date": self.publication_date,
            "original_url": self.original_url,
            "domain": self.domain,
            "credibility_info": self.credibility_info,
            **self.metadata,
        }


def _clean_text(text: str) -> str:
    return " ".join(text.split())


def _safe_date(value: Any) -> Optional[str]:
    if value in (None, "", "unknown"):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        try:
            datetime.fromisoformat(value[:10])
            return value[:10]
        except ValueError:
            return value
    return str(value)


def parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def parse_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def parse_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def parse_csv(path: Path) -> List[Dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="", errors="ignore") as file:
        return list(csv.DictReader(file))


def parse_json(path: Path) -> List[Dict[str, Any]]:
    with path.open("r", encoding="utf-8", errors="ignore") as file:
        data = json.load(file)
    if isinstance(data, dict):
        if "documents" in data and isinstance(data["documents"], list):
            return data["documents"]
        return [data]
    if isinstance(data, list):
        return data
    return []


def load_document_file(path: Path) -> List[SourceDocument]:
    suffix = path.suffix.lower()
    docs: List[SourceDocument] = []

    if suffix == ".txt":
        text = parse_txt(path)
        docs.append(
            SourceDocument(
                document_id=path.stem,
                title=path.stem.replace("_", " ").title(),
                source=str(path),
                source_type="local_text",
                domain="unknown",
                text=_clean_text(text),
            )
        )

    elif suffix == ".pdf":
        text = parse_pdf(path)
        docs.append(
            SourceDocument(
                document_id=path.stem,
                title=path.stem.replace("_", " ").title(),
                source=str(path),
                source_type="pdf",
                domain="unknown",
                text=_clean_text(text),
            )
        )

    elif suffix == ".docx":
        text = parse_docx(path)
        docs.append(
            SourceDocument(
                document_id=path.stem,
                title=path.stem.replace("_", " ").title(),
                source=str(path),
                source_type="docx",
                domain="unknown",
                text=_clean_text(text),
            )
        )

    elif suffix == ".csv":
        rows = parse_csv(path)
        for index, row in enumerate(rows):
            text = row.get("text") or row.get("content") or " ".join(str(v) for v in row.values() if v)
            docs.append(
                SourceDocument(
                    document_id=row.get("document_id") or f"{path.stem}_{index}",
                    title=row.get("title") or f"{path.stem}_{index}",
                    source=row.get("source") or str(path),
                    source_type=row.get("source_type") or "csv",
                    publication_date=_safe_date(row.get("publication_date")),
                    credibility_info={key: value for key, value in row.items() if key.startswith("cred") or key in {"reliability", "domain"}},
                    original_url=row.get("url") or row.get("original_url") or "",
                    domain=row.get("domain") or "unknown",
                    text=_clean_text(text),
                    metadata={k: v for k, v in row.items() if k not in {"text", "content", "document_id", "title", "source", "source_type", "publication_date", "url", "original_url", "domain"}},
                )
            )

    elif suffix == ".json":
        rows = parse_json(path)
        for index, row in enumerate(rows):
            text = row.get("text") or row.get("content") or row.get("body") or " "
            docs.append(
                SourceDocument(
                    document_id=row.get("document_id") or row.get("id") or f"{path.stem}_{index}",
                    title=row.get("title") or row.get("name") or f"{path.stem}_{index}",
                    source=row.get("source") or row.get("publisher") or str(path),
                    source_type=row.get("source_type") or row.get("kind") or "json",
                    publication_date=_safe_date(row.get("publication_date") or row.get("date")),
                    credibility_info=row.get("credibility_info") or {},
                    original_url=row.get("original_url") or row.get("url") or "",
                    domain=row.get("domain") or "unknown",
                    text=_clean_text(text),
                    metadata={k: v for k, v in row.items() if k not in {"text", "content", "body", "document_id", "id", "title", "name", "source", "publisher", "source_type", "kind", "publication_date", "date", "original_url", "url", "domain", "credibility_info"}},
                )
            )

    return docs


def load_documents_from_directory(directory: str | Path) -> List[SourceDocument]:
    base_dir = Path(directory)
    documents: List[SourceDocument] = []
    if not base_dir.exists():
        return documents

    for path in sorted(base_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".txt", ".pdf", ".docx", ".csv", ".json"}:
            documents.extend(load_document_file(path))
    return documents


def load_all_documents(paths: Iterable[str | Path]) -> List[SourceDocument]:
    all_docs: List[SourceDocument] = []
    for path in paths:
        base = Path(path)
        if base.is_dir():
            all_docs.extend(load_documents_from_directory(base))
        elif base.is_file():
            all_docs.extend(load_document_file(base))
    return all_docs
